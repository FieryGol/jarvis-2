# plugins/fitxers.py
# Buscar, llegir, crear, copiar, moure i esborrar fitxers i carpetes.
# Esborrar envia a la Paperera (es pot recuperar). Moure, esborrar i
# sobreescriure demanen permís a l'usuari.

import os
import shutil
import time
from pathlib import Path

from eines import boolea, eina, enter, text

DESCRIPCIO = "Fitxers i carpetes de l'ordinador"

INICI = Path.home()


def _carpeta_coneguda(nom):
    """'Escriptori', 'Documents', 'Baixades'... → la ruta real (també si és a OneDrive)."""
    nom = nom.strip().lower()
    noms = {
        "escriptori": "Desktop", "desktop": "Desktop", "escritorio": "Desktop",
        "documents": "Documents", "documentos": "Documents",
        "baixades": "Downloads", "descàrregues": "Downloads", "descargas": "Downloads", "downloads": "Downloads",
        "imatges": "Pictures", "imágenes": "Pictures", "fotos": "Pictures", "pictures": "Pictures",
        "música": "Music", "musica": "Music", "music": "Music",
        "vídeos": "Videos", "videos": "Videos",
        "usuari": "", "inici": "", "home": "",
    }
    if nom not in noms:
        return None
    sub = noms[nom]
    if not sub:
        return INICI
    for base in (INICI / "OneDrive", INICI):
        if (base / sub).exists():
            return base / sub
    return INICI / sub


def ruta(valor):
    """Converteix el que diu l'usuari en una ruta. Accepta 'Escriptori/treball.docx'."""
    valor = os.path.expandvars(valor.strip().strip('"'))
    p = Path(valor)
    if p.is_absolute():
        return p
    primera, *resta = p.parts or [""]
    coneguda = _carpeta_coneguda(primera)
    if coneguda is not None:
        return coneguda.joinpath(*resta)
    return INICI / p


def _mida(n):
    for unitat in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.0f} {unitat}"
        n /= 1024
    return f"{n:.1f} TB"


@eina("Busca fitxers o carpetes pel nom (o part del nom) a les carpetes de l'usuari. "
      "Fes-la servir abans d'obrir, llegir o moure un fitxer si no en saps la ruta exacta.",
      parametres={"nom": text("Nom o part del nom, per exemple 'apunts física' o '.pdf'"),
                  "carpeta": text("On buscar: Escriptori, Documents, Baixades... (per defecte totes les de l'usuari)")})
def buscar_fitxers(nom, carpeta=None):
    base = ruta(carpeta) if carpeta else INICI
    paraules = nom.lower().split()
    trobats = []
    limit_temps = time.time() + 8
    salta = {"appdata", "node_modules", ".git", "__pycache__", "$recycle.bin", ".cache", "venv"}
    for arrel, carpetes, fitxers in os.walk(base):
        carpetes[:] = [c for c in carpetes if c.lower() not in salta and not c.startswith(".")]
        for f in carpetes + fitxers:
            if all(p in f.lower() for p in paraules):
                trobats.append(Path(arrel) / f)
        if len(trobats) >= 30 or time.time() > limit_temps:
            break
    if not trobats:
        return f"No he trobat res que es digui «{nom}» a {base}."
    trobats.sort(key=lambda p: -p.stat().st_mtime if p.exists() else 0)
    return "He trobat:\n" + "\n".join(str(p) for p in trobats[:20])


@eina("Llista el que hi ha dins d'una carpeta.",
      parametres={"carpeta": text("Ruta o nom de la carpeta (Escriptori, Baixades, C:\\...)")})
def llistar_carpeta(carpeta):
    p = ruta(carpeta)
    if not p.is_dir():
        return f"No existeix la carpeta {p}."
    elements = sorted(p.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower()))
    linies = [f"{'[carpeta] ' if e.is_dir() else ''}{e.name}"
              + ("" if e.is_dir() else f" ({_mida(e.stat().st_size)})") for e in elements[:60]]
    mes = f"\n... i {len(elements) - 60} més" if len(elements) > 60 else ""
    return f"{p}:\n" + "\n".join(linies) + mes


@eina("Llegeix el contingut d'un fitxer de text, Word (.docx) o PDF per poder-lo resumir, "
      "corregir o respondre preguntes sobre ell.",
      parametres={"fitxer": text("Ruta del fitxer")})
def llegir_fitxer(fitxer):
    p = ruta(fitxer)
    if not p.is_file():
        return f"No existeix el fitxer {p}."
    ext = p.suffix.lower()
    try:
        if ext == ".docx":
            import docx
            contingut = "\n".join(par.text for par in docx.Document(p).paragraphs)
        elif ext == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                return "Per llegir PDF cal instal·lar la biblioteca pypdf (pip install pypdf)."
            contingut = "\n".join((pg.extract_text() or "") for pg in PdfReader(p).pages[:30])
        else:
            contingut = p.read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        return f"No he pogut llegir {p.name}: {e}"
    if len(contingut) > 12000:
        contingut = contingut[:12000] + "\n[... el fitxer continua, només n'he llegit el principi]"
    return contingut or "El fitxer és buit."


@eina("Obre un fitxer o una carpeta amb el programa que toca (un PDF amb el lector, "
      "una foto amb el visor, una carpeta amb l'explorador...).",
      parametres={"fitxer": text("Ruta del fitxer o carpeta")})
def obrir_fitxer(fitxer):
    p = ruta(fitxer)
    if not p.exists():
        return f"No existeix {p}. Prova primer amb buscar_fitxers."
    os.startfile(p)
    return f"He obert {p.name}."


@eina("Crea un fitxer de text amb el contingut indicat (per exemple una llista, un resum o "
      "un codi). Si ja existeix, només el sobreescriu si sobreescriure és cert.",
      parametres={"fitxer": text("Ruta del fitxer, per exemple Escriptori/llista.txt"),
                  "contingut": text("El text que ha de tenir"),
                  "sobreescriure": boolea("Si ja existeix, substituir-lo")},
      requerits=["fitxer", "contingut"],
      perillosa=lambda a: bool(a.get("sobreescriure")) and ruta(a.get("fitxer", "")).exists(),
      resum=lambda a: f"Sobreescriure el fitxer {ruta(a.get('fitxer', ''))}")
def crear_fitxer(fitxer, contingut, sobreescriure=False):
    p = ruta(fitxer)
    if p.exists() and not sobreescriure:
        return f"Ja existeix {p}. Pregunta a l'usuari si el vol substituir."
    p.parent.mkdir(parents=True, exist_ok=True)
    if p.suffix.lower() == ".docx":
        import docx
        d = docx.Document()
        for linia in contingut.split("\n"):
            d.add_paragraph(linia)
        d.save(p)
    else:
        p.write_text(contingut, encoding="utf-8")
    return f"He creat {p}."


@eina("Crea una carpeta nova.", parametres={"carpeta": text("Ruta de la carpeta nova")})
def crear_carpeta(carpeta):
    p = ruta(carpeta)
    p.mkdir(parents=True, exist_ok=True)
    return f"Carpeta a punt: {p}."


@eina("Copia un fitxer o una carpeta a un altre lloc.",
      parametres={"origen": text("Què s'ha de copiar"), "desti": text("Carpeta o ruta de destí")})
def copiar(origen, desti):
    o, d = ruta(origen), ruta(desti)
    if not o.exists():
        return f"No existeix {o}."
    if d.is_dir():
        d = d / o.name
    if d.exists():
        return f"Ja hi ha {d}; no el sobreescric."
    if o.is_dir():
        shutil.copytree(o, d)
    else:
        d.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(o, d)
    return f"Copiat a {d}."


@eina("Mou o canvia el nom d'un fitxer o una carpeta.",
      parametres={"origen": text("Què s'ha de moure"), "desti": text("Carpeta de destí o nom nou")},
      perillosa=True, resum=lambda a: f"Moure {ruta(a.get('origen', ''))} → {ruta(a.get('desti', ''))}")
def moure(origen, desti):
    o, d = ruta(origen), ruta(desti)
    if not o.exists():
        return f"No existeix {o}."
    if d.is_dir():
        d = d / o.name
    if d.exists():
        return f"Ja hi ha {d}; no el sobreescric."
    shutil.move(str(o), str(d))
    return f"Mogut a {d}."


@eina("Envia un fitxer o una carpeta a la Paperera de reciclatge (es pot recuperar).",
      parametres={"fitxer": text("Ruta del fitxer o carpeta")},
      perillosa=True, resum=lambda a: f"Enviar a la Paperera: {ruta(a.get('fitxer', ''))}")
def esborrar(fitxer):
    p = ruta(fitxer)
    if not p.exists():
        return f"No existeix {p}."
    from send2trash import send2trash
    send2trash(str(p))
    return f"He enviat {p.name} a la Paperera."


@eina("Diu quant ocupa una carpeta i quins són els fitxers més grans que hi ha.",
      parametres={"carpeta": text("Ruta de la carpeta"), "quants": enter("Quants fitxers grans mostrar (per defecte 10)")},
      requerits=["carpeta"])
def fitxers_grans(carpeta, quants=10):
    p = ruta(carpeta)
    total, tots = 0, []
    limit_temps = time.time() + 10
    for arrel, _, fitxers in os.walk(p):
        for f in fitxers:
            try:
                mida = (Path(arrel) / f).stat().st_size
            except OSError:
                continue
            total += mida
            tots.append((mida, Path(arrel) / f))
        if time.time() > limit_temps:
            break
    tots.sort(reverse=True)
    return (f"{p} ocupa uns {_mida(total)}. Els més grans:\n"
            + "\n".join(f"{_mida(m)}  {f}" for m, f in tots[:int(quants)]))
